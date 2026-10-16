"""Independent Decimal fixtures and editable cumulative PAYE workbook formulas.
The calendar-month annualisation method is an implementation assumption for review,
not a formula prescribed verbatim by the linked JRB guidance or accountant signoff.
"""
from decimal import Decimal as D, ROUND_HALF_UP
from pathlib import Path
import calendar,json
from openpyxl.styles import Font,PatternFill
round_money=lambda x:x.quantize(D('.01'),rounding=ROUND_HALF_UP)
kobo=lambda x:int(round_money(x)*100)

def add_cumulative_review(w):
 inputs=[
  ('January standard','2026-01','2026-01-01',dict(basic='200000',pension=False),{},dict(gross='0',paye='0')),
  ('October standard','2026-10','2026-01-01',dict(basic='200000',pension=False),{},dict(gross='1800000',paye='180000')),
  ('Bonus added once','2026-10','2026-01-01',dict(basic='200000',pension=False),dict(bonus='100000'),dict(gross='1800000',paye='180000')),
  ('Salary increase','2026-10','2026-01-01',dict(basic='300000',pension=False),{},dict(gross='1800000',paye='180000')),
  ('Pension NHF rent','2026-10','2026-01-01',dict(basic='200000',housing='100000',transport='50000',nhf=True,annual_rent='1200000'),{},dict(gross='3150000',paye='323640',pension='252000',nhf='45000')),
  ('First job October starter','2026-10','2026-10-16',dict(basic='310000',pension=False),{},dict(gross='0',paye='0')),
  ('December annual reconciliation','2026-12','2026-01-01',dict(basic='250000',pension=False),{},dict(gross='2750000',paye='302500')),
  ('December top band','2026-12','2026-01-01',dict(basic='5000000',pension=False),{},dict(gross='55000000',paye='11852500')),
  ('Kobo rounding','2026-07','2026-01-01',dict(basic='200000.01',nhis='12.34',annual_rent='987654.32'),{},dict(gross='1200000.06',paye='60000',pension='96000',nhis='74.04')),
  ('Confirmed minimum wage exemption','2026-10','2026-01-01',dict(basic='70000',pension=False),dict(minimum_wage_exempt=True),dict(gross='630000',paye='0')),
  ('Prior under-deduction recovered','2026-10','2026-01-01',dict(basic='200000'),{},dict(gross='1800000',pension='144000',paye='90000')),
 ]
 cases=[];sheet=w.create_sheet('Cumulative PAYE')
 sheet.append(['Case','Calendar month','Prior gross','Prior pension','Prior NHF','Prior NHIS','Prior PAYE','Current gross','Current pension','Current NHF','Current NHIS','Current other deductions','Annual rent','Other annual reliefs','Min wage exempt?','Annual rent relief','Annualised taxable','Raw annual band tax','Cumulative tax','Current PAYE','Total current deductions','Current net','Expected PAYE','Expected net','PAYE delta','Net delta'])
 for i,(name,period,joined,raw,a,p) in enumerate(inputs,2):
  month=int(period[5:]);days=calendar.monthrange(2026,month)[1];paid=days if joined[:7]<period else days-int(joined[8:])+1
  e={k:round_money(D(raw.get(k,'0'))*paid/days) for k in ['basic','housing','transport','other']}
  gross=sum(e.values())+D(a.get('bonus','0'));pension=round_money((e['basic']+e['housing']+e['transport']+D(a.get('pensionable_bonus','0')))*D('.08')) if raw.get('pension',True) else D(0)
  employer=round_money((e['basic']+e['housing']+e['transport']+D(a.get('pensionable_bonus','0')))*D('.10')) if raw.get('pension',True) else D(0)
  nhf=round_money(e['basic']*D('.025')) if raw.get('nhf',False) else D(0);nhis=D(raw.get('nhis','0'));other=D(raw.get('other_deductions','0'))+D(a.get('one_off_deduction','0'))
  prior={k:D(p.get(k,'0')) for k in ['gross','paye','pension','nhf','nhis']}
  rent=min(round_money(D(raw.get('annual_rent','0'))*D('.2')),D(500000));reliefs=D(raw.get('annual_reliefs','0'))
  taxable=max(D(0),(prior['gross']+gross-prior['pension']-pension-prior['nhf']-nhf-prior['nhis']-nhis)*12/month-rent-reliefs)
  tax=sum(max(D(0),min(taxable,upper)-lower)*rate for lower,upper,rate in zip(map(D,[0,800000,3000000,12000000,25000000,50000000]),map(D,[800000,3000000,12000000,25000000,50000000,'1e20']),map(D,['0','.15','.18','.21','.23','.25'])))
  cumulative=D(0) if a.get('minimum_wage_exempt',False) else round_money(tax*month/12);paye=max(D(0),cumulative-prior['paye']);deductions=paye+pension+nhf+nhis+other;net=gross-deductions
  assert net>=0 and cumulative>=prior['paye']
  adjustment={**a,'paye_method':'cumulative','history_confirmed':True,'review_note':'Engineering method fixture; accountant approval pending'}
  context=dict(through_period='2025-12' if month==1 else f'2026-{month-1:02d}',totals={k:kobo(v) for k,v in prior.items()},unrecorded_periods=[])
  expected={k:kobo(v) for k,v in dict(gross=gross,paye=paye,pension=pension,employer_pension=employer,nhf=nhf,nhis=nhis,total_deductions=deductions,net=net,employer_ecs=round_money(gross*D('.01')),employer_itf=D(0)).items()}
  cases.append(dict(name=name,period=period,joined=joined,input=raw,adjustment=adjustment,prior=context,expected=expected,expected_cumulative_tax=kobo(cumulative)))
  sheet.append([name,month]+[float(prior[k]) for k in ['gross','pension','nhf','nhis','paye']]+[float(v) for v in [gross,pension,nhf,nhis,other,D(raw.get('annual_rent','0')),reliefs]]+['Yes' if a.get('minimum_wage_exempt',False) else 'No'])
  formulas={'P':f'=MIN(ROUND(M{i}*0.2,2),500000)','Q':f'=MAX(0,(C{i}+H{i}-D{i}-I{i}-E{i}-J{i}-F{i}-K{i})*12/B{i}-P{i}-N{i})','R':f'=MIN(MAX(Q{i}-800000,0),2200000)*0.15+MIN(MAX(Q{i}-3000000,0),9000000)*0.18+MIN(MAX(Q{i}-12000000,0),13000000)*0.21+MIN(MAX(Q{i}-25000000,0),25000000)*0.23+MAX(Q{i}-50000000,0)*0.25','S':f'=IF(O{i}="Yes",0,ROUND(R{i}*B{i}/12,2))','T':f'=MAX(0,S{i}-G{i})','U':f'=SUM(I{i}:L{i})+T{i}','V':f'=H{i}-U{i}','Y':f'=T{i}-W{i}','Z':f'=V{i}-X{i}'}
  for col,formula in formulas.items():sheet[f'{col}{i}']=formula
  sheet[f'W{i}']=float(paye);sheet[f'X{i}']=float(net)
  for cell in sheet[i][2:]:
   if cell.column!=15:cell.number_format='#,##0.00'
  for cell in sheet[i][1:15]:cell.font=Font(color='1565C0');cell.fill=PatternFill('solid',fgColor='EDF5FF')
 assert cases[0]['expected']['paye']==2000000
 assert cases[2]['expected']['paye']==3500000
 assert cases[6]['expected']['paye']==2750000
 assert cases[7]['expected']['paye']==107750000
 sheet.column_dimensions['A'].width=40
 from openpyxl.utils import get_column_letter
 for col in range(2,27):sheet.column_dimensions[get_column_letter(col)].width=23
 sheet.row_dimensions[1].height=50;sheet.auto_filter.ref=sheet.dimensions
 sheet.append([]);sheet.append(['Method assumption','Calendar-month YTD annualisation with proportional annual reliefs/bands. The JRB source requires cumulative records but does not prescribe this formula verbatim. Confirm method with the relevant accountant/tax authority.'])
 sheet.append(['Eligibility','All prior months must be covered. Eligible pension/NHF/NHIS contributions and evidence must be reviewed. Minimum-wage exemption requires explicit evidence confirmation.'])
 sheet.append(['Blocked cases','Negative current tax indicates a possible refund; negative net indicates excessive recovery. App preview flags these and blocks automatic drafting. Workbook does not perform refunds.'])
 Path('public/payroll-cumulative-cases.json').write_text(json.dumps(cases,indent=2)+'\n')
 w['Sources'].append(['JRB PIT Guidelines 2026','https://irs.kg.gov.ng/JRB_2026_PIT_Guidelines_Final.pdf','Sections 7–9, 13 and Appendix 1: cumulative records, eligible deductions, rent relief, non-periodic payments and 2026 bands. Calendar-month apportionment is an explicit implementation assumption.'])
 return cases
