"""Independent Decimal reference calculations and an editable Excel review workbook.
Not an accountant approval. Run from the repository root; openpyxl is required.
"""
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.properties import CalcProperties
D=Decimal
round_money=lambda n:n.quantize(D('.01'),rounding=ROUND_HALF_UP)
kobo=lambda n:int(round_money(n)*100)
inputs=[
 ('Minimum wage',dict(basic='70000',pension=False)),
 ('Just above minimum wage',dict(basic='70000.01',pension=False)),
 ('Basic salary with pension',dict(basic='200000')),
 ('PAYE band: annual 3m',dict(basic='250000',pension=False)),
 ('PAYE band: annual 12m',dict(basic='1000000',pension=False)),
 ('PAYE band: around annual 25m',dict(basic='2083333.33',pension=False)),
 ('PAYE band: around annual 50m',dict(basic='4166666.67',pension=False)),
 ('PAYE top band: annual 60m',dict(basic='5000000',pension=False)),
 ('Pension NHF and rent relief',dict(basic='200000',housing='100000',transport='50000',nhf=True,annual_rent='1200000')),
 ('Rent relief cap',dict(basic='350000',annual_rent='4000000')),
 ('NHIS and other deductions',dict(basic='250000',nhis='10000',other_deductions='5000')),
 ('Other recurring earnings',dict(basic='200000',housing='50000',transport='25000',other='75000',nhf=True)),
 ('Other eligible annual reliefs',dict(basic='200000',annual_reliefs='120000')),
 ('Kobo and rounding',dict(basic='123456.78',housing='23456.79',transport='3456.78',nhf=True,annual_rent='987654.32',nhis='1234.56',other_deductions='111.11')),
 ('Taxable zero after reliefs',dict(basic='100000',annual_reliefs='1200000',pension=False)),
]
fields=['basic','housing','transport','other','annual_rent','nhis','other_deductions','annual_reliefs']
def reference(raw):
 s={k:D(str(raw.get(k,0))) for k in fields}
 gross=sum(s[k] for k in ['basic','housing','transport','other'])
 pension=round_money((s['basic']+s['housing']+s['transport'])*D('.08')) if raw.get('pension',True) else D(0)
 employer=round_money((s['basic']+s['housing']+s['transport'])*D('.10')) if raw.get('pension',True) else D(0)
 nhf=round_money(s['basic']*D('.025')) if raw.get('nhf',False) else D(0)
 rent=min(round_money(s['annual_rent']*D('.20')),D(500000))
 taxable=max(D(0),(gross-pension-nhf-s['nhis'])*12-rent-s['annual_reliefs'])
 cumulative_limits=[D(800000),D(3000000),D(12000000),D(25000000),D(50000000)]
 tax=D(0)
 # Independent cumulative threshold expression, rather than the engine's remaining-width loop.
 for lower,upper,rate in zip([D(0)]+cumulative_limits,cumulative_limits+[D('1e20')],[D(0),D('.15'),D('.18'),D('.21'),D('.23'),D('.25')]):
  tax+=max(D(0),min(taxable,upper)-lower)*rate
 paye=D(0) if gross<=70000 else round_money(tax/12)
 deductions=pension+nhf+s['nhis']+s['other_deductions']+paye
 return {k:kobo(v) for k,v in dict(gross=gross,paye=paye,pension=pension,employer_pension=employer,nhf=nhf,nhis=s['nhis'],total_deductions=deductions,net=gross-deductions,employer_ecs=round_money(gross*D('.01')),employer_itf=D(0)).items()}
cases=[dict(name=name,input=raw,expected=reference(raw)) for name,raw in inputs]
# Explicit hand-worked anchors keep reference generation from silently drifting.
assert cases[3]['expected']['paye']==2750000
assert cases[8]['expected']['paye']==3596000 and cases[8]['expected']['net']==28104000
assert cases[7]['expected']['paye']==107750000
Path('public/payroll-reference-cases.json').write_text(json.dumps(cases,indent=2)+'\n')
# Variable periods deliberately take reviewed PAYE rather than estimating cumulative tax.
period_inputs=[
 ('October starter and bonus','2026-10','2026-10-16',dict(basic='310000',pension=False),dict(bonus='50000',one_off_deduction='5000',reviewed_paye='12000',review_note='Engineering fixture: supplied PAYE; not accountant signoff')),
 ('Starter pensionable bonus NHF','2026-10','2026-10-16',dict(basic='310000',nhf=True),dict(bonus='50000',pensionable_bonus='20000',one_off_deduction='5000',reviewed_paye='8000',review_note='Engineering fixture: supplied PAYE')),
 ('February half month','2026-02','2026-02-15',dict(basic='200000'),dict(reviewed_paye='3000',review_note='Engineering fixture: supplied PAYE')),
 ('Full month pensionable bonus','2026-10','2026-01-01',dict(basic='200000'),dict(bonus='100000',pensionable_bonus='100000',reviewed_paye='25000',review_note='Engineering fixture: supplied PAYE')),
 ('Partial minimum wage explicit zero','2026-10','2026-10-31',dict(basic='70000',pension=False),dict(reviewed_paye='0',review_note='Engineering fixture: supplied zero PAYE')),
 ('Component rounding','2026-10','2026-10-17',dict(basic='123456.78',housing='23456.79',transport='3456.78',other='111.11',nhf=True,nhis='123.45',other_deductions='11.11'),dict(bonus='500.01',pensionable_bonus='100.01',one_off_deduction='2.22',reviewed_paye='1500.01',review_note='Engineering fixture: rounding')),
]
period_cases=[]
import calendar
for name,period,joined,raw,a in period_inputs:
 days=calendar.monthrange(2026,int(period[5:]))[1];paid=days if joined[:7]<period else days-int(joined[8:])+1
 e={k:round_money(D(raw.get(k,'0'))*paid/days) for k in ['basic','housing','transport','other']}
 bonus=D(a.get('bonus','0'));gross=sum(e.values())+bonus;base=sum(e[k] for k in ['basic','housing','transport'])+D(a.get('pensionable_bonus','0'))
 pension=round_money(base*D('.08')) if raw.get('pension',True) else D(0);employer=round_money(base*D('.10')) if raw.get('pension',True) else D(0)
 nhf=round_money(e['basic']*D('.025')) if raw.get('nhf',False) else D(0);nhis=D(raw.get('nhis','0'));paye=D(a['reviewed_paye']);deduct=pension+nhf+nhis+paye+D(raw.get('other_deductions','0'))+D(a.get('one_off_deduction','0'))
 expected={k:kobo(v) for k,v in dict(gross=gross,paye=paye,pension=pension,employer_pension=employer,nhf=nhf,nhis=nhis,total_deductions=deduct,net=gross-deduct,employer_ecs=round_money(gross*D('.01')),employer_itf=D(0)).items()}
 period_cases.append(dict(name=name,period=period,joined=joined,input=raw,adjustment=a,expected=expected,paid_days=paid,days=days))
assert period_cases[0]['expected']['net']==19300000
assert period_cases[1]['expected']['net']==17860000
Path('public/payroll-period-cases.json').write_text(json.dumps(period_cases,indent=2)+'\n')
w=Workbook();readme=w.active;readme.title='Read me'
notes=[['Nigeria 2026 payroll calculation review'],['Status','Engineering reference cases — NOT accountant-approved'],['Scope','Recurring monthly earnings plus reviewed period adjustments; Nigeria 2026 only.'],['Instructions','Change blue inputs on Calculation. Review assumptions and evidence. Reconcile every case with your accountant.'],['Rounding','Positive half-up to two decimal naira places (integer kobo). Monthly PAYE rounds the unrounded annual band tax divided by 12.'],['Pension','8% employee / 10% employer on basic + housing + transport; confirm contractual pensionable emoluments and applicability.'],['NHF','Opt-in/applicability per employee. Default 2.5% basic. Official guidance differs on basis; confirm before approving.'],['PAYE reliefs','Pension, NHF, NHIS, declared rent relief (20%, capped 500,000), and approved other eligible annual reliefs.'],['Minimum wage','Monthly gross at or below 70,000 treated as PAYE-exempt; confirm the employee’s legal eligibility.'],['Employer levies','ECS default 1% gross. ITF is disabled until applicability is confirmed; normally 1% annual payroll for eligible employers.'],['Period adjustments','Calendar-day starters, bonuses and one-off deductions. Period review sheet uses supplied reviewed PAYE, not a cumulative tax estimate. NHIS and recurring after-tax deductions are not prorated.'],['Not supported','Automatic cumulative/YTD PAYE, arrears, leaver calculations, unpaid absence, foreign currency or nonresident cases.'],['Sources','See Sources sheet. Tax legislation and accountant signoff control over this workbook.'],['Independent expected values','Reference sheet uses Python Decimal with cumulative thresholds; Calculation uses editable Excel marginal-band formulas.'],['YTD reconciliation','Separate sheet verifies opening + released totals and draft exclusion. This does not compute cumulative PAYE liability.'],['Signoff','Accountant: __________________ Date: __________ Findings: __________________'],['Launch target','Company signup → first approved payroll <=24 hours. This is measured in app analytics; it is not a guarantee.']]
for r in notes:readme.append(r)
readme.column_dimensions['A'].width=28;readme.column_dimensions['B'].width=112
for row in readme:
 for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
readme.row_dimensions[1].height=30
for n in range(2,len(notes)+1):readme.row_dimensions[n].height=40
settings=w.create_sheet('Settings')
for row in [['Parameter','Value'],['Employee pension',.08],['Employer pension',.10],['NHF rate',.025],['NHF basis','basic'],['Rent relief rate',.20],['Rent relief maximum',500000],['Minimum wage',70000],['Employer ECS rate',.01],['Employer ITF rate',0],['Pension basis','bht']]:settings.append(row)
settings.column_dimensions['A'].width=32;settings.column_dimensions['B'].width=20
calc=w.create_sheet('Calculation');headers=['Case','Basic','Housing','Transport','Other earnings','Annual rent','NHIS monthly','Other deductions','Annual other reliefs','Pension?','NHF?','Gross','Pension','Employer pension','NHF','Rent relief','Annual taxable','Raw annual band tax','Monthly PAYE','Total deductions','Net pay','Employer ECS','Employer ITF','Expected PAYE','Expected net','PAYE delta','Net delta']
calc.append(headers)
reference_sheet=w.create_sheet('Reference');reference_sheet.append(['Case']+list(cases[0]['expected']))
for i,case in enumerate(cases,2):
 raw=case['input'];calc.append([case['name']]+[float(D(str(raw.get(k,0)))) for k in fields]+['Yes' if raw.get('pension',True) else 'No','Yes' if raw.get('nhf',False) else 'No'])
 formulas={
 'L':f'=SUM(B{i}:E{i})','M':f'=IF(J{i}="Yes",ROUND(IF(Settings!$B$11="gross",L{i},SUM(B{i}:D{i}))*Settings!$B$2,2),0)',
 'N':f'=IF(J{i}="Yes",ROUND(IF(Settings!$B$11="gross",L{i},SUM(B{i}:D{i}))*Settings!$B$3,2),0)',
 'O':f'=IF(K{i}="Yes",ROUND(IF(Settings!$B$5="gross",L{i},B{i})*Settings!$B$4,2),0)',
 'P':f'=MIN(ROUND(F{i}*Settings!$B$6,2),Settings!$B$7)',
 'Q':f'=MAX(0,(L{i}-M{i}-O{i}-G{i})*12-P{i}-I{i})',
 'R':f'=MIN(MAX(Q{i}-800000,0),2200000)*0.15+MIN(MAX(Q{i}-3000000,0),9000000)*0.18+MIN(MAX(Q{i}-12000000,0),13000000)*0.21+MIN(MAX(Q{i}-25000000,0),25000000)*0.23+MAX(Q{i}-50000000,0)*0.25',
 'S':f'=IF(L{i}<=Settings!$B$8,0,ROUND(R{i}/12,2))','T':f'=SUM(M{i},O{i},G{i},H{i},S{i})','U':f'=L{i}-T{i}',
 'V':f'=ROUND(L{i}*Settings!$B$9,2)','W':f'=ROUND(L{i}*Settings!$B$10,2)','Z':f'=S{i}-X{i}','AA':f'=U{i}-Y{i}'}
 for col,formula in formulas.items():calc[f'{col}{i}']=formula
 calc[f'X{i}']=case['expected']['paye']/100;calc[f'Y{i}']=case['expected']['net']/100
 reference_sheet.append([case['name']]+[v/100 for v in case['expected'].values()])
 for row in calc.iter_rows(min_row=i,max_row=i,min_col=2,max_col=9):
  for c in row:c.font=Font(color='1565C0')
 for c in calc[i][1:]:
  if c.column not in (10,11):c.number_format='#,##0.00;[Red](#,##0.00)'
validation=DataValidation(type='list',formula1='"Yes,No"');calc.add_data_validation(validation);validation.add(f'J2:K{len(cases)+1}')
calc.freeze_panes='B2';calc.auto_filter.ref=calc.dimensions
calc.column_dimensions['A'].width=35
for col in range(2,28):calc.column_dimensions[__import__('openpyxl').utils.get_column_letter(col)].width=20
sources=w.create_sheet('Sources')
for row in [['Source','URL','Use / review note'],['Nigeria Tax Act 2025','https://www.nrs.gov.ng/uploads/NIGERIA_TAX_ACT_2025_ef6bb812a5.pdf','2026 bands and eligible deductions; legislation controls.'],['NRS legislation index','https://www.nrs.gov.ng/tax-laws/page','Official law index.'],['Plateau State IRS 2026 calculator','https://www.psirs.gov.ng/tax-calculator','Effective date and rent-relief guidance.'],['PenCom 2023 FAQ','https://www.pencom.gov.ng/wp-content/uploads/2023/09/FREQUENTLY-ASKED-QUESTIONS-2023.pdf','Monthly emoluments; 8% employee / 10% employer minimum.'],['FMBN current FAQ','https://fmbn.gov.ng/products/nhf-scheme/faqs','Current wording refers to gross income. Confirm NHF base.'],['FMBN service charter','https://www.fmbn.gov.ng/documents/FMBN_SERVICOM_CHARTER.pdf','Earlier wording refers to basic salary; confirm applicable guidance.'],['NSITF services','https://nsitf.gov.ng/services','Employer ECS 1% total monthly payroll.'],['ITF','https://www.itf.gov.ng/','Eligibility and 1% annual payroll contribution.']]:sources.append(row)
for n,width in [('A',32),('B',90),('C',85)]:sources.column_dimensions[n].width=width
for row in sources:
 for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
period_sheet=w.create_sheet('Period review')
period_sheet.append(['Case','Calendar days','Paid days','Basic monthly','Housing monthly','Transport monthly','Other monthly','Bonus','Pensionable bonus BHT','Reviewed PAYE INPUT','NHIS fixed','Recurring deductions fixed','One-off deduction','Pension?','NHF?','Gross','Employee pension','Employer pension','NHF','Total deductions','Net','Employer ECS','Employer ITF','Expected net','Net delta'])
for i,c in enumerate(period_cases,2):
 raw=c['input'];a=c['adjustment']
 period_sheet.append([c['name'],c['days'],c['paid_days']]+[float(D(raw.get(k,'0'))) for k in ['basic','housing','transport','other']]+[float(D(a.get(k,'0'))) for k in ['bonus','pensionable_bonus','reviewed_paye']]+[float(D(raw.get(k,'0'))) for k in ['nhis','other_deductions']]+[float(D(a.get('one_off_deduction','0')))]+['Yes' if raw.get('pension',True) else 'No','Yes' if raw.get('nhf',False) else 'No'])
 components=[f'ROUND({col}{i}*C{i}/B{i},2)' for col in 'DEFG']
 gross='+'.join(components)+f'+H{i}'
 bht='+'.join(components[:3])+f'+I{i}'
 formulas={'P':f'={gross}','Q':f'=IF(N{i}="Yes",ROUND(IF(Settings!$B$11="gross",P{i},{bht})*Settings!$B$2,2),0)','R':f'=IF(N{i}="Yes",ROUND(IF(Settings!$B$11="gross",P{i},{bht})*Settings!$B$3,2),0)','S':f'=IF(O{i}="Yes",ROUND(IF(Settings!$B$5="gross",P{i},{components[0]})*Settings!$B$4,2),0)','T':f'=SUM(J{i}:M{i})+Q{i}+S{i}','U':f'=P{i}-T{i}','V':f'=ROUND(P{i}*Settings!$B$9,2)','W':f'=ROUND(P{i}*Settings!$B$10,2)','Y':f'=U{i}-X{i}'}
 for col,formula in formulas.items():period_sheet[f'{col}{i}']=formula
 period_sheet[f'X{i}']=c['expected']['net']/100
 for cell in period_sheet[i][1:]:
  if cell.column not in (2,3,14,15):cell.number_format='#,##0.00'
 for cell in period_sheet[i][1:15]:cell.font=Font(color='1565C0');cell.fill=PatternFill('solid',fgColor='EDF5FF')
period_sheet.column_dimensions['A'].width=38
for col in range(2,26):period_sheet.column_dimensions[__import__('openpyxl').utils.get_column_letter(col)].width=23
period_sheet.auto_filter.ref=period_sheet.dimensions
period_sheet.row_dimensions[1].height=45
# YTD reconciliation uses arithmetic, not a new tax-liability formula.
ytd=w.create_sheet('YTD reconciliation')
ytd.append(['Scenario','Opening gross','Opening PAYE','Opening employee pension','Opening other employee deductions','Released gross','Released PAYE','Released employee pension','Released other employee deductions','Draft gross EXCLUDED','YTD gross','YTD PAYE withheld','YTD deductions','YTD net','Expected gross','Expected net','Gross delta','Net delta'])
ytd_cases=[
 ('Opening only',1800000,90000,144000,0,0,0,0,0,200000,1800000,1566000),
 ('October released',1800000,90000,144000,0,200000,17600,16000,0,0,2000000,1732400),
 ('November draft excluded',1800000,90000,144000,0,200000,17600,16000,0,200000,2000000,1732400),
 ('October and November released',1800000,90000,144000,0,400000,35200,32000,0,0,2200000,1898800),
]
for i,case in enumerate(ytd_cases,2):
 name,*values=case;ytd.append([name]+values[:9])
 for col,formula in {'K':f'=B{i}+F{i}','L':f'=C{i}+G{i}','M':f'=SUM(C{i}:E{i})+SUM(G{i}:I{i})','N':f'=K{i}-M{i}','Q':f'=K{i}-O{i}','R':f'=N{i}-P{i}'}.items():ytd[f'{col}{i}']=formula
 ytd[f'O{i}']=values[9];ytd[f'P{i}']=values[10]
 for cell in ytd[i][1:]:cell.number_format='#,##0.00'
 for cell in ytd[i][1:10]:cell.font=Font(color='1565C0');cell.fill=PatternFill('solid',fgColor='EDF5FF')
ytd.column_dimensions['A'].width=40
for col in range(2,19):ytd.column_dimensions[__import__('openpyxl').utils.get_column_letter(col)].width=25
ytd.row_dimensions[1].height=50
ytd.append([]);ytd.append(['Basis','Opening balances cover January through a reviewed cutoff; only Approved/Paid payroll is included afterwards. Drafts do not affect totals.'])
ytd.append(['Tax status','PAYE is withheld, not confirmed remittance. This sheet does not calculate cumulative tax liability or verify external accountant identity.'])
ytd.append(['Evidence','Opening references and change history are company-admin only. Released payroll review records include a calculation SHA-256 fingerprint.'])
for sheet in w:
 sheet.freeze_panes='A2'
 for cell in sheet[1]:cell.fill=PatternFill('solid',fgColor='245B45');cell.font=Font(color='FFFFFF',bold=True);cell.alignment=Alignment(wrap_text=True)
 sheet.sheet_view.showGridLines=False
 if sheet==reference_sheet:
  sheet.column_dimensions['A'].width=40
  for col in range(2,12):sheet.column_dimensions[__import__('openpyxl').utils.get_column_letter(col)].width=24
calc.row_dimensions[1].height=44
reference_sheet.row_dimensions[1].height=38
for row in calc.iter_rows(min_row=2,min_col=2,max_col=11):
 for cell in row:
  cell.font=Font(color='1565C0');cell.fill=PatternFill('solid',fgColor='EDF5FF')
for cell in ['B2','B3','B4','B6','B9','B10']:settings[cell].number_format='0.00%'
for cell in ['B7','B8']:settings[cell].number_format='#,##0.00'
for row in sources.iter_rows(min_row=2):
 row[1].hyperlink=row[1].value;row[1].style='Hyperlink'
w.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True)
w.save('public/payroll-review.xlsx')
print(f'Created review workbook and {len(cases)} independent reference cases. Not accountant-approved.')
